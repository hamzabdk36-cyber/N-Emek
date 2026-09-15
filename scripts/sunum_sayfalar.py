"""Jüri sunumunun icerik sayfalari: md okuyucu, cizim yardimcilari, yerlesimler.

`sunum_pptx.py` sablonu hazirlar (sayfa sirasi, yonergeler, kapak, takim,
alt bant) ve her icerik sayfasi icin buradaki `YERLESIMLER` tablosundan
ilgili fonksiyonu cagirir.

Ayrim bilincli: **metin** `docs/SUNUM/NN-*.md` icinde, **yerlesim** burada.
Bir sayfanin metni degisince kod degismez; bir kutunun yeri degisince md
degismez. Metin tasarsa `sunum_powerpoint.ps1` sayfa ve kutu adiyla soyler.

md bicimi
---------
Dosya `## anahtar` basliklariyla bolumlere ayrilir; yerlesim fonksiyonu
bolumleri anahtarla ister. Eksik ya da artik bolum hata verir - sessizce
bos kutu cizmektense durmak iyi.

Bolum icinde:

    ### Alt baslik
    - madde
    Duz paragraf (bos satirla ayrilir)
    > vurgulu satir
    | tablo | satiri |
    ![aciklama](../gorseller/x.jpg){kirp=sol,ust,sag,alt}

Satir ici `**kalin**`. `<!-- ... -->` yorumlari atilir (`sayim:`
isaretleri icin; `dokuman_denetimi.py` md dosyasini okur).

Iki ozel bolum: `## not` konusmaci notuna, `## kaynak` sayfanin altindaki
kaynak satirina gider.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Cm, Pt
from PIL import Image

ONEK = "NEMEK"

# Sablon paleti + projenin renk disiplini (CLAUDE.md "Arayuz renk
# disiplini"): altin para ve pay, yesil dogrulanmis, kirmizi itiraz ve
# dusuk guven, mavi zincir ve baglanti.
LACIVERT = RGBColor(0x1A, 0x42, 0x6A)
METIN = RGBColor(0x1F, 0x2A, 0x37)
GRI = RGBColor(0x5B, 0x66, 0x75)
GRI_ACIK = RGBColor(0xD5, 0xDB, 0xE3)
ZEMIN = RGBColor(0xEE, 0xF3, 0xF9)
MAVI_ACIK = RGBColor(0xB7, 0xCC, 0xE4)
MAVI = RGBColor(0x2F, 0x6F, 0xB0)
TURUNCU = RGBColor(0xF3, 0xA0, 0x25)
ALTIN = RGBColor(0xA8, 0x6A, 0x0C)
YESIL = RGBColor(0x2E, 0x7D, 0x4F)
KIRMIZI = RGBColor(0xB0, 0x3A, 0x2E)
BEYAZ = RGBColor(0xFF, 0xFF, 0xFF)

SOL, SAG, ALT = 3.4, 41.0, 25.2
GEN = SAG - SOL
KAYNAK_UST = 24.2

YAZI = "Arial"


# --------------------------------------------------------------------------
# md okuyucu
# --------------------------------------------------------------------------

@dataclass
class Blok:
    tur: str                      # baslik | madde | paragraf | vurgu | tablo | gorsel
    metin: str = ""
    satirlar: list[list[str]] = field(default_factory=list)
    yol: str = ""
    kirp: tuple[int, int, int, int] | None = None


def bolumleri_oku(yol: Path) -> dict[str, list[Blok] | str]:
    ham = re.sub(r"<!--.*?-->", "", yol.read_text(encoding="utf-8"), flags=re.S)
    bolumler: dict[str, list[str]] = {}
    anahtar = None
    for satir in ham.splitlines():
        m = re.match(r"^##\s+(\S+)\s*$", satir)
        if m:
            anahtar = m.group(1)
            if anahtar in bolumler:
                raise SystemExit(f"{yol.name}: '{anahtar}' bölümü iki kez var")
            bolumler[anahtar] = []
        elif anahtar is not None:
            bolumler[anahtar].append(satir.rstrip())
    sonuc: dict[str, list[Blok] | str] = {}
    for ad, satirlar in bolumler.items():
        if ad in ("not", "kaynak"):
            sonuc[ad] = "\n".join(s for s in satirlar).strip()
        else:
            sonuc[ad] = bloklara_ayir(satirlar)
    return sonuc


def bloklara_ayir(satirlar: list[str]) -> list[Blok]:
    bloklar: list[Blok] = []
    paragraf: list[str] = []

    def paragrafi_kapat() -> None:
        if paragraf:
            bloklar.append(Blok("paragraf", " ".join(paragraf)))
            paragraf.clear()

    for satir in satirlar:
        s = satir.strip()
        if not s:
            paragrafi_kapat()
            continue
        if s.startswith("### "):
            paragrafi_kapat()
            bloklar.append(Blok("baslik", s[4:].strip()))
        elif s.startswith("- "):
            paragrafi_kapat()
            bloklar.append(Blok("madde", s[2:].strip()))
        elif s.startswith("> "):
            paragrafi_kapat()
            bloklar.append(Blok("vurgu", s[2:].strip()))
        elif s.startswith("|"):
            paragrafi_kapat()
            hucreler = [h.strip() for h in s.strip("|").split("|")]
            if all(re.fullmatch(r":?-{3,}:?", h) for h in hucreler):
                continue
            if bloklar and bloklar[-1].tur == "tablo":
                bloklar[-1].satirlar.append(hucreler)
            else:
                bloklar.append(Blok("tablo", satirlar=[hucreler]))
        elif s.startswith("!["):
            paragrafi_kapat()
            m = re.fullmatch(r"!\[(.*?)\]\((.+?)\)(?:\{kirp=(\d+),(\d+),(\d+),(\d+)\})?", s)
            if not m:
                raise SystemExit(f"görsel satırı okunamadı: {s}")
            kirp = tuple(int(x) for x in m.groups()[2:]) if m.group(3) else None
            bloklar.append(Blok("gorsel", m.group(1), yol=m.group(2), kirp=kirp))
        else:
            paragraf.append(s)
    paragrafi_kapat()
    return bloklar


class Icerik:
    """Bir sayfanin bolumleri; kullanilmayan bolum kalirsa hata verir."""

    def __init__(self, yol: Path):
        self.yol = yol
        self.bolumler = bolumleri_oku(yol)
        self.kullanilan: set[str] = set()

    def __getitem__(self, anahtar: str) -> list[Blok]:
        if anahtar not in self.bolumler:
            raise SystemExit(f"{self.yol.name}: '## {anahtar}' bölümü yok")
        self.kullanilan.add(anahtar)
        return self.bolumler[anahtar]  # type: ignore[return-value]

    def varsa(self, anahtar: str):
        if anahtar in self.bolumler:
            self.kullanilan.add(anahtar)
            return self.bolumler[anahtar]
        return None

    def maddeler(self, anahtar: str) -> list[str]:
        return [b.metin for b in self[anahtar] if b.tur == "madde"]

    def tablo(self, anahtar: str) -> list[list[str]]:
        tablolar = [b for b in self[anahtar] if b.tur == "tablo"]
        if len(tablolar) != 1:
            raise SystemExit(f"{self.yol.name}: '{anahtar}' bölümünde tek tablo bekleniyordu")
        return tablolar[0].satirlar

    def gorsel(self, anahtar: str) -> Blok:
        gorseller = [b for b in self[anahtar] if b.tur == "gorsel"]
        if len(gorseller) != 1:
            raise SystemExit(f"{self.yol.name}: '{anahtar}' bölümünde tek görsel bekleniyordu")
        return gorseller[0]

    def artiklari_denetle(self) -> None:
        artik = set(self.bolumler) - self.kullanilan - {"not"}
        if artik:
            raise SystemExit(f"{self.yol.name}: kullanılmayan bölüm: {sorted(artik)}")


def ayir(madde: str) -> tuple[str, str]:
    """`**Baslik** · aciklama` -> (Baslik, aciklama)."""
    m = re.fullmatch(r"\*\*(.+?)\*\*\s*(?:·\s*)?(.*)", madde)
    return (m.group(1), m.group(2)) if m else (madde, "")


# --------------------------------------------------------------------------
# Cizim yardimcilari
# --------------------------------------------------------------------------

def _adla(sekil, ad: str, alt_sinir: float | None = None) -> None:
    sekil.name = f"{ONEK} {ad}" + (f" @{alt_sinir:.2f}" if alt_sinir is not None else "")


def kosular(paragraf, metin: str, boyut: float, renk: RGBColor, kalin: bool = False) -> None:
    for parca in re.split(r"(\*\*.+?\*\*)", metin):
        if not parca:
            continue
        vurgulu = parca.startswith("**") and parca.endswith("**")
        run = paragraf.add_run()
        run.text = parca[2:-2] if vurgulu else parca
        run.font.size = Pt(boyut)
        run.font.name = YAZI
        run.font.bold = kalin or vurgulu
        run.font.color.rgb = renk


def _madde_isareti(paragraf, boyut: float, renk: RGBColor) -> None:
    ppr = paragraf._p.get_or_add_pPr()
    girinti = int(Pt(boyut) * 0.95)
    ppr.set("marL", str(girinti))
    ppr.set("indent", str(-girinti))
    for eski in ppr.findall(qn("a:buNone")):
        ppr.remove(eski)
    renk_el = ppr.makeelement(qn("a:buClr"), {})
    renk_el.append(renk_el.makeelement(qn("a:srgbClr"), {"val": str(renk)}))
    ppr.append(renk_el)
    ppr.append(ppr.makeelement(qn("a:buFont"), {"typeface": YAZI}))
    ppr.append(ppr.makeelement(qn("a:buChar"), {"char": "•"}))


def cerceveyi_hazirla(cerceve, ic: float = 0.0, dikey=MSO_ANCHOR.TOP) -> None:
    cerceve.word_wrap = True
    cerceve.auto_size = MSO_AUTO_SIZE.NONE
    cerceve.vertical_anchor = dikey
    cerceve.margin_left = cerceve.margin_right = Cm(ic)
    cerceve.margin_top = cerceve.margin_bottom = Cm(ic * 0.7)


def bloklari_yaz(cerceve, bloklar: list[Blok], boyut: float, renk: RGBColor = METIN,
                 baslik_renk: RGBColor = LACIVERT, hizala=None) -> None:
    ilk = True
    for blok in bloklar:
        if blok.tur in ("tablo", "gorsel"):
            raise SystemExit("metin kutusunda tablo/görsel olamaz; ayrı bölüme taşıyın")
        p = cerceve.paragraphs[0] if ilk else cerceve.add_paragraph()
        if hizala is not None:
            p.alignment = hizala
        if blok.tur == "baslik":
            p.space_before = Pt(0 if ilk else boyut * 0.7)
            kosular(p, blok.metin, boyut + 2, baslik_renk, kalin=True)
        elif blok.tur == "madde":
            p.space_before = Pt(0 if ilk else boyut * 0.35)
            _madde_isareti(p, boyut, baslik_renk)
            kosular(p, blok.metin, boyut, renk)
        elif blok.tur == "vurgu":
            p.space_before = Pt(0 if ilk else boyut * 0.6)
            kosular(p, blok.metin, boyut, baslik_renk, kalin=True)
        else:
            p.space_before = Pt(0 if ilk else boyut * 0.5)
            kosular(p, blok.metin, boyut, renk)
        ilk = False


def metin(slayt, ad: str, x: float, y: float, w: float, h: float, bloklar: list[Blok],
          boyut: float = 18, renk: RGBColor = METIN, baslik_renk: RGBColor = LACIVERT,
          hizala=None, dikey=MSO_ANCHOR.TOP):
    kutu = slayt.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
    _adla(kutu, ad)
    cerceveyi_hazirla(kutu.text_frame, 0.0, dikey)
    bloklari_yaz(kutu.text_frame, bloklar, boyut, renk, baslik_renk, hizala)
    return kutu


def kart(slayt, ad: str, x: float, y: float, w: float, h: float, bloklar: list[Blok],
         boyut: float = 16, dolgu: RGBColor = ZEMIN, renk: RGBColor = METIN,
         baslik_renk: RGBColor = LACIVERT, cizgi: RGBColor | None = None, ic: float = 0.45,
         hizala=None, dikey=MSO_ANCHOR.TOP, sol_serit: RGBColor | None = None):
    sekil = slayt.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Cm(x), Cm(y), Cm(w), Cm(h))
    sekil.adjustments[0] = min(0.12, 0.35 / min(w, h))
    _adla(sekil, ad)
    sekil.shadow.inherit = False
    sekil.fill.solid()
    sekil.fill.fore_color.rgb = dolgu
    if cizgi is None:
        sekil.line.fill.background()
    else:
        sekil.line.color.rgb = cizgi
        sekil.line.width = Pt(1)
    cerceveyi_hazirla(sekil.text_frame, ic, dikey)
    # Otomatik sekillerin varsayilan hizasi ortali; kartlarda metin sola dayali.
    bloklari_yaz(sekil.text_frame, bloklar, boyut, renk, baslik_renk,
                 PP_ALIGN.LEFT if hizala is None else hizala)
    if sol_serit is not None:
        serit = slayt.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(x), Cm(y + 0.25), Cm(0.18), Cm(h - 0.5))
        _adla(serit, f"{ad}-serit")
        serit.shadow.inherit = False
        serit.fill.solid()
        serit.fill.fore_color.rgb = sol_serit
        serit.line.fill.background()
    return sekil


def sayi_karti(slayt, ad: str, x: float, y: float, w: float, h: float, buyuk: str, aciklama: str,
               renk: RGBColor = LACIVERT, buyuk_boyut: float = 32, boyut: float = 13,
               dolgu: RGBColor = ZEMIN):
    sekil = kart(slayt, ad, x, y, w, h, [], boyut, dolgu=dolgu, ic=0.35, dikey=MSO_ANCHOR.MIDDLE)
    cer = sekil.text_frame
    p = cer.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    kosular(p, buyuk, buyuk_boyut, renk, kalin=True)
    p2 = cer.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    p2.space_before = Pt(2)
    kosular(p2, aciklama, boyut, GRI)
    return sekil


def ok(slayt, ad: str, x: float, y: float, w: float, h: float, renk: RGBColor = MAVI,
       yon: str = "sag"):
    tur = {"sag": MSO_SHAPE.RIGHT_ARROW, "asagi": MSO_SHAPE.DOWN_ARROW}[yon]
    sekil = slayt.shapes.add_shape(tur, Cm(x), Cm(y), Cm(w), Cm(h))
    _adla(sekil, ad)
    sekil.shadow.inherit = False
    sekil.fill.solid()
    sekil.fill.fore_color.rgb = renk
    sekil.line.fill.background()
    return sekil


def akis(slayt, ad: str, x: float, y: float, w: float, h: float, ogeler: list[str],
         boyut: float = 16, alt_boyut: float = 12.5, ok_gen: float = 0.8,
         dolgular: list[RGBColor] | None = None, yazi_renkleri: list[RGBColor] | None = None):
    """Yatay akis: `**Baslik** · aciklama` maddelerinden kutular ve oklar."""
    n = len(ogeler)
    kutu_w = (w - (n - 1) * ok_gen) / n
    for i, madde in enumerate(ogeler):
        baslik, aciklama = ayir(madde)
        kx = x + i * (kutu_w + ok_gen)
        dolgu = dolgular[i] if dolgular else ZEMIN
        yazi = yazi_renkleri[i] if yazi_renkleri else LACIVERT
        bloklar = [Blok("baslik", baslik)] + ([Blok("paragraf", aciklama)] if aciklama else [])
        kart(slayt, f"{ad}-{i + 1}", kx, y, kutu_w, h, bloklar, alt_boyut, dolgu=dolgu,
             renk=yazi if dolgular else METIN, baslik_renk=yazi, ic=0.3,
             hizala=PP_ALIGN.CENTER, dikey=MSO_ANCHOR.MIDDLE)
        # baslik puntosu: bloklari_yaz basliga +2 veriyor; akista ayri ayar
        for run in slayt.shapes[-1].text_frame.paragraphs[0].runs:
            run.font.size = Pt(boyut)
        if i < n - 1:
            ok(slayt, f"{ad}-ok-{i + 1}", kx + kutu_w + ok_gen * 0.18, y + h / 2 - 0.35,
               ok_gen * 0.64, 0.7, MAVI)


def tablo(slayt, ad: str, x: float, y: float, w: float, satirlar: list[list[str]],
          oranlar: list[float], boyut: float = 14, satir_h: float = 1.0, alt_sinir: float = ALT,
          hizalar: str | None = None, vurgu_son: bool = False, renkli_hucreler: bool = False):
    """Ilk satir baslik. `hizalar`: sutun basina 'l', 'c' veya 'r'."""
    n_satir, n_sutun = len(satirlar), len(satirlar[0])
    if any(len(s) != n_sutun for s in satirlar):
        raise SystemExit(f"tablo '{ad}': satırların sütun sayısı eşit değil")
    gf = slayt.shapes.add_table(n_satir, n_sutun, Cm(x), Cm(y), Cm(w), Cm(satir_h * n_satir))
    _adla(gf, f"tablo-{ad}", alt_sinir)
    tbl = gf.table
    tbl.first_row = True
    tbl.horz_banding = False
    toplam = sum(oranlar)
    for j, oran in enumerate(oranlar):
        tbl.columns[j].width = Cm(w * oran / toplam)
    for i, satir in enumerate(satirlar):
        tbl.rows[i].height = Cm(min(satir_h, 1.15) if i == 0 else satir_h)
        son = vurgu_son and i == n_satir - 1
        for j, deger in enumerate(satir):
            hucre = tbl.cell(i, j)
            hucre.margin_left = hucre.margin_right = Cm(0.22)
            hucre.margin_top = hucre.margin_bottom = Cm(0.08)
            hucre.vertical_anchor = MSO_ANCHOR.MIDDLE
            hucre.fill.solid()
            if i == 0:
                hucre.fill.fore_color.rgb = LACIVERT
            elif son:
                hucre.fill.fore_color.rgb = MAVI_ACIK
            else:
                hucre.fill.fore_color.rgb = ZEMIN if i % 2 else BEYAZ
            cer = hucre.text_frame
            cer.word_wrap = True
            p = cer.paragraphs[0]
            hiz = (hizalar or "l" * n_sutun)[j]
            p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[hiz]
            renk = BEYAZ if i == 0 else METIN
            if renkli_hucreler and i > 0:
                renk = {"var": YESIL, "kısmen": ALTIN, "yok": KIRMIZI}.get(deger, renk)
            kosular(p, deger, boyut, renk, kalin=(i == 0 or son or (renkli_hucreler and i > 0 and j > 0)))
    return gf


def gorsel(slayt, ad: str, x: float, y: float, w: float, h: float, blok: Blok, kok: Path,
           hizala: str = "orta", cerceve: bool = True):
    yol = (kok / blok.yol).resolve()
    if not yol.exists():
        raise SystemExit(f"görsel yok: {yol}")
    with Image.open(yol) as im:
        gw, gh = im.size
    sl, su, sa, sal = blok.kirp or (0, 0, gw, gh)
    kirpik_w, kirpik_h = sa - sl, sal - su
    olcek = min(w / kirpik_w, h / kirpik_h)
    cw, ch = kirpik_w * olcek, kirpik_h * olcek
    cx = x + (w - cw) / 2 if hizala == "orta" else x
    cy = y + (h - ch) / 2
    pic = slayt.shapes.add_picture(str(yol), Cm(cx), Cm(cy), Cm(cw), Cm(ch))
    _adla(pic, f"gorsel-{ad}")
    pic.crop_left, pic.crop_top = sl / gw, su / gh
    pic.crop_right, pic.crop_bottom = 1 - sa / gw, 1 - sal / gh
    pic._element.xpath("./p:nvPicPr/p:cNvPr")[0].set("descr", blok.metin)
    if cerceve:
        pic.line.color.rgb = GRI_ACIK
        pic.line.width = Pt(0.75)
    return pic, (cx, cy, cw, ch)


def kaynak_satiri(slayt, metin_: str) -> None:
    if not metin_:
        return
    kutu = slayt.shapes.add_textbox(Cm(SOL), Cm(KAYNAK_UST), Cm(GEN), Cm(ALT - KAYNAK_UST))
    _adla(kutu, "kaynak")
    cerceveyi_hazirla(kutu.text_frame, 0.0, MSO_ANCHOR.BOTTOM)
    p = kutu.text_frame.paragraphs[0]
    kosular(p, "Kaynak: " + " ".join(metin_.split()), 11, GRI)


def notu_yaz(slayt, metin_: str | None) -> None:
    if metin_:
        slayt.notes_slide.notes_text_frame.text = " ".join(metin_.split())


def B(*satirlar: str) -> list[Blok]:
    """Kod icinde kucuk etiket bloklari (yerlesimin parcasi, icerik degil)."""
    return [Blok("baslik", s) for s in satirlar]


# --------------------------------------------------------------------------
# Sayfa yerlesimleri. Her biri (slayt, icerik, ust, kok) alir.
# --------------------------------------------------------------------------

def s_ozet(sl, c: Icerik, ust: float, kok: Path) -> None:
    metin(sl, "ana", SOL, ust + 0.2, GEN, 2.6, c["ana"], boyut=24, renk=LACIVERT)
    akis(sl, "akis", SOL, ust + 3.3, GEN, 4.3, c.maddeler("akis"), boyut=17, alt_boyut=13.5)
    ust2 = ust + 8.4
    w = (GEN - 2 * 0.7) / 3
    for i, anahtar in enumerate(("konu", "yontem", "tema")):
        kart(sl, anahtar, SOL + i * (w + 0.7), ust2, w, ALT - ust2, c[anahtar], boyut=19,
             sol_serit=TURUNCU, ic=0.6)


def s_problem(sl, c: Icerik, ust: float, kok: Path) -> None:
    sw = 18.3
    sagx = SOL + sw + 0.9
    saw = SAG - sagx
    # sol: zincir, bilesenler, hukuk
    metin(sl, "zincir-baslik", SOL, ust + 0.1, sw, 0.9, B("Bir paylaşım zinciri nasıl kopar"), boyut=16)
    akis(sl, "zincir", SOL, ust + 1.1, sw, 3.9, c.maddeler("zincir"), boyut=16, alt_boyut=12.5,
         ok_gen=0.7, dolgular=[ZEMIN, ZEMIN, RGBColor(0xFB, 0xEC, 0xE9)],
         yazi_renkleri=[LACIVERT, LACIVERT, KIRMIZI])
    metin(sl, "bilesenler", SOL, ust + 5.5, sw, 9.6, c["bilesenler"], boyut=17)
    kart(sl, "hukuk", SOL, ust + 15.4, sw, KAYNAK_UST - 0.3 - (ust + 15.4), c["hukuk"], boyut=14,
         sol_serit=LACIVERT)
    # sag: istatistik + tablo
    istat = c.maddeler("istatistik")
    kw = (saw - 0.6) / 2
    for i, madde in enumerate(istat):
        buyuk, aciklama = ayir(madde)
        sayi_karti(sl, f"istatistik-{i + 1}", sagx + i * (kw + 0.6), ust + 0.1, kw, 4.9,
                   buyuk, aciklama, renk=ALTIN, buyuk_boyut=30, boyut=12.5)
    metin(sl, "tablo-baslik", sagx, ust + 5.5, saw, 0.9, B("Mevcut yaklaşımlar ve eksik kaldıkları yer"), boyut=16)
    tablo(sl, "mevcut", sagx, ust + 6.5, saw, c.tablo("mevcut"), [0.9, 1.25], boyut=13.5,
          satir_h=1.45, alt_sinir=KAYNAK_UST - 0.2)
    kaynak_satiri(sl, c.bolumler.get("kaynak", ""))  # type: ignore[arg-type]
    c.kullanilan.add("kaynak")


def s_cozum_1(sl, c: Icerik, ust: float, kok: Path) -> None:
    metin(sl, "ana", SOL, ust + 0.2, GEN, 1.8, c["ana"], boyut=24, renk=LACIVERT)
    sw = 19.6
    y0 = ust + 2.5
    _, (gx, gy, gw, gh) = gorsel(sl, "olculen-bolge", SOL, y0, sw, 9.6, c.gorsel("bolge_gorsel"), kok)
    metin(sl, "bolge-aciklama", SOL, gy + gh + 0.3, sw, 1.6, c["bolge_aciklama"], boyut=13, renk=GRI)
    metin(sl, "olcum", SOL, gy + gh + 2.1, sw, KAYNAK_UST - 0.2 - (gy + gh + 2.1), c["olcum"], boyut=15)
    sagx = SOL + sw + 0.9
    saw = SAG - sagx
    kart(sl, "formul", sagx, y0, saw, 2.2, c["formul"], boyut=20, dolgu=LACIVERT, renk=BEYAZ,
         baslik_renk=BEYAZ, hizala=PP_ALIGN.CENTER, dikey=MSO_ANCHOR.MIDDLE)
    _, (_, ky, _, kh) = gorsel(sl, "emek-karti", sagx, y0 + 2.6, saw, 5.2, c.gorsel("kart_gorsel"), kok)
    metin(sl, "carpanlar", sagx, ky + kh + 0.4, saw, KAYNAK_UST - 0.2 - (ky + kh + 0.4), c["carpanlar"],
          boyut=15)
    kaynak_satiri(sl, c.bolumler.get("kaynak", ""))  # type: ignore[arg-type]
    c.kullanilan.add("kaynak")


def s_cozum_2(sl, c: Icerik, ust: float, kok: Path) -> None:
    metin(sl, "ana", SOL, ust + 0.2, GEN, 1.8, c["ana"], boyut=24, renk=LACIVERT)
    sw = 23.2
    y0 = ust + 2.5
    metin(sl, "matris-baslik", SOL, y0, sw, 0.9, B("Yetenek karşılaştırması"), boyut=16)
    tablo(sl, "matris", SOL, y0 + 1.0, sw, c.tablo("matris"), [1.55, 1, 1, 1, 1, 1], boyut=13,
          satir_h=1.25, hizalar="lccccc", vurgu_son=True, renkli_hucreler=True, alt_sinir=y0 + 10.6)
    metin(sl, "matris-not", SOL, y0 + 10.0, sw, 1.3, c["matris_not"], boyut=11.5, renk=GRI)
    metin(sl, "itiraz-baslik", SOL, y0 + 11.8, sw, 0.9, B("İtiraz akışı"), boyut=16)
    akis(sl, "itiraz", SOL, y0 + 12.8, sw, ALT - (y0 + 12.8), c.maddeler("itiraz"), boyut=14,
         alt_boyut=12, ok_gen=0.6)
    sagx = SOL + sw + 0.9
    saw = SAG - sagx
    kart(sl, "asla", sagx, y0, saw, 9.9, c["asla"], boyut=16.5, dolgu=LACIVERT, renk=BEYAZ,
         baslik_renk=BEYAZ, ic=0.55)
    kart(sl, "katki", sagx, y0 + 10.4, saw, ALT - (y0 + 10.4), c["katki"], boyut=15.5,
         sol_serit=TURUNCU)


def s_yontem_1(sl, c: Icerik, ust: float, kok: Path) -> None:
    sw = 19.4
    y0 = ust + 0.1
    metin(sl, "mimari-baslik", SOL, y0, sw, 0.9, B("Sistem mimarisi"), boyut=16)
    katmanlar = c.maddeler("mimari")
    motor = c.maddeler("motor")
    # katman sirasi: istemci, api, servisler, [motor: iki kutu], depolama
    y = y0 + 1.1
    kat_h, ara = 2.55, 0.75
    alt_y = KAYNAK_UST - 0.3
    kat_h = (alt_y - y - 4 * ara) / 5
    satirlar = katmanlar[:3] + [None] + katmanlar[3:]
    for i, madde in enumerate(satirlar):
        if madde is None:
            yw = (sw - 0.6) / 2
            for k, m in enumerate(motor):
                baslik, aciklama = ayir(m)
                kart(sl, f"motor-{k + 1}", SOL + k * (yw + 0.6), y, yw, kat_h,
                     [Blok("baslik", baslik), Blok("paragraf", aciklama)], boyut=12.5,
                     dolgu=LACIVERT, renk=BEYAZ, baslik_renk=BEYAZ, ic=0.3,
                     hizala=PP_ALIGN.CENTER, dikey=MSO_ANCHOR.MIDDLE)
        else:
            baslik, aciklama = ayir(madde)
            kart(sl, f"katman-{i + 1}", SOL, y, sw, kat_h,
                 [Blok("baslik", baslik), Blok("paragraf", aciklama)], boyut=12.5, ic=0.3,
                 hizala=PP_ALIGN.CENTER, dikey=MSO_ANCHOR.MIDDLE)
        if i < len(satirlar) - 1:
            ok(sl, f"mimari-ok-{i + 1}", SOL + sw / 2 - 0.3, y + kat_h + 0.08, 0.6, ara - 0.16, MAVI, "asagi")
        y += kat_h + ara
    sagx = SOL + sw + 0.9
    saw = SAG - sagx
    metin(sl, "hat-baslik", sagx, y0, saw, 0.9, B("Altı aşamalı köken kurtarma hattı"), boyut=16)
    tablo(sl, "asamalar", sagx, y0 + 1.1, saw, c.tablo("asamalar"), [0.3, 1.25, 2.45, 0.95], boyut=12.5,
          satir_h=1.05, hizalar="clll", alt_sinir=y0 + 9.1)
    metin(sl, "hat-not", sagx, y0 + 9.25, saw, 1.4, c["hat_not"], boyut=12.5, renk=GRI)
    kart(sl, "yz", sagx, y0 + 10.9, saw, KAYNAK_UST - 0.3 - (y0 + 10.9), c["yz"], boyut=14,
         sol_serit=TURUNCU)
    kaynak_satiri(sl, c.bolumler.get("kaynak", ""))  # type: ignore[arg-type]
    c.kullanilan.add("kaynak")


def s_yontem_2(sl, c: Icerik, ust: float, kok: Path) -> None:
    y0 = ust + 0.1
    sw = 18.3
    sagx = SOL + sw + 0.9
    saw = SAG - sagx
    kart(sl, "veri", SOL, y0, sw, 8.6, c["veri"], boyut=14.5, sol_serit=MAVI)
    metrikler = c.maddeler("metrikler")
    kw, kh = (saw - 0.6) / 2, (8.6 - 0.6) / 2
    renkler = [YESIL, YESIL, YESIL, LACIVERT]
    for i, madde in enumerate(metrikler):
        buyuk, aciklama = ayir(madde)
        sayi_karti(sl, f"metrik-{i + 1}", sagx + (i % 2) * (kw + 0.6), y0 + (i // 2) * (kh + 0.6),
                   kw, kh, buyuk, aciklama, renk=renkler[i], buyuk_boyut=28, boyut=12)
    y1 = y0 + 9.2
    kart(sl, "indirgeme", SOL, y1, sw, 5.2, c["indirgeme"], boyut=13.5, sol_serit=MAVI)
    kart(sl, "ozel-kapsama", sagx, y1, saw, 5.2, c["ozel_kapsama"], boyut=13.5, sol_serit=MAVI)
    y2 = y1 + 5.8
    etik = c["etik"]
    basliklar = [b for b in etik if b.tur == "baslik"]
    maddeler = [b for b in etik if b.tur == "madde"]
    yari = (len(maddeler) + 1) // 2
    metin(sl, "etik-baslik", SOL, y2, GEN, 0.9, basliklar, boyut=14)
    kh2 = KAYNAK_UST - 0.3 - (y2 + 1.0)
    kart(sl, "etik-1", SOL, y2 + 1.0, sw, kh2, maddeler[:yari], boyut=13, sol_serit=YESIL, ic=0.35)
    kart(sl, "etik-2", sagx, y2 + 1.0, saw, kh2, maddeler[yari:], boyut=13, sol_serit=YESIL, ic=0.35)
    kaynak_satiri(sl, c.bolumler.get("kaynak", ""))  # type: ignore[arg-type]
    c.kullanilan.add("kaynak")


def s_prototip(sl, c: Icerik, ust: float, kok: Path) -> None:
    y0 = ust + 0.1
    akis(sl, "gelistirme", SOL, y0, GEN, 2.7, c.maddeler("gelistirme"), boyut=15, alt_boyut=12,
         ok_gen=0.7, dolgular=[RGBColor(0xE4, 0xF1, 0xE9)] * 4, yazi_renkleri=[YESIL] * 4)
    y1 = y0 + 3.2
    gh = 8.2
    gorseller = [c.gorsel("ekran_akis"), c.gorsel("ekran_bolge"), c.gorsel("ekran_mobil")]
    # Uc gorsel ayni yukseklikte, en-boy oranlarina gore; toplam genislik
    # alani asarsa hepsi birlikte kucultulur, bosluklar esit kalir.
    oranlar = []
    for blok in gorseller:
        with Image.open((kok / blok.yol).resolve()) as im:
            gw0, gh0 = im.size
        sl_, su_, sa_, sal_ = blok.kirp or (0, 0, gw0, gh0)
        oranlar.append((sa_ - sl_) / (sal_ - su_))
    ara_g = 0.8
    h_g = min(gh, (GEN - ara_g * (len(oranlar) - 1)) / sum(oranlar))
    toplam = h_g * sum(oranlar) + ara_g * (len(oranlar) - 1)
    x = SOL + (GEN - toplam) / 2
    for i, (blok, oran) in enumerate(zip(gorseller, oranlar)):
        gorsel(sl, f"ekran-{i + 1}", x, y1 + (gh - h_g) / 2, h_g * oran, h_g, blok, kok)
        x += h_g * oran + ara_g
    metin(sl, "ekran-aciklama", SOL, y1 + gh + 0.15, GEN, 0.8, c["ekran_aciklama"], boyut=12, renk=GRI,
          hizala=PP_ALIGN.CENTER)
    y2 = y1 + gh + 1.1
    w = (GEN - 2 * 0.6) / 3
    for i, anahtar in enumerate(("tasarim", "erisilebilirlik", "kullanici")):
        kart(sl, anahtar, SOL + i * (w + 0.6), y2, w, ALT - y2, c[anahtar], boyut=12.5, ic=0.35,
             sol_serit=TURUNCU)


def s_uygulanabilirlik_1(sl, c: Icerik, ust: float, kok: Path) -> None:
    y0 = ust + 0.1
    sw = 21.4
    kart(sl, "model", SOL, y0, sw, 4.0, c["model"], boyut=15, sol_serit=ALTIN)
    metin(sl, "dagitim-baslik", SOL, y0 + 4.4, sw, 0.9, B("Aynı ₺50.000 havuz, iki model"), boyut=16)
    tablo(sl, "dagitim", SOL, y0 + 5.4, sw, c.tablo("dagitim"), [1.7, 1.2, 1.2, 0.7], boyut=15,
          satir_h=1.1, hizalar="lrrr", alt_sinir=y0 + 11.1)
    metin(sl, "dagitim-not", SOL, y0 + 11.1, sw, 1.3, c["dagitim_not"], boyut=11.5, renk=GRI)
    kart(sl, "havuz", SOL, y0 + 12.6, sw, KAYNAK_UST - 0.3 - (y0 + 12.6), c["havuz"], boyut=16,
         dolgu=LACIVERT, renk=BEYAZ, baslik_renk=BEYAZ, dikey=MSO_ANCHOR.MIDDLE)
    sagx = SOL + sw + 0.9
    saw = SAG - sagx
    kart(sl, "ticari", sagx, y0, saw, 8.0, c["ticari"], boyut=15, sol_serit=ALTIN)
    metin(sl, "benimseme-baslik", sagx, y0 + 8.5, saw, 0.9, B("Benimseme yolu"), boyut=16)
    adimlar = c.maddeler("benimseme")
    yb = y0 + 9.5
    ah = (KAYNAK_UST - 0.3 - yb - (len(adimlar) - 1) * 0.3) / len(adimlar)
    for i, madde in enumerate(adimlar):
        baslik, aciklama = ayir(madde)
        kutu = kart(sl, f"benimseme-{i + 1}", sagx, yb + i * (ah + 0.3), saw, ah,
                    [Blok("paragraf", f"**{i + 1} · {baslik}** · {aciklama}")], boyut=12.5, ic=0.3,
                    dikey=MSO_ANCHOR.MIDDLE, sol_serit=MAVI)
    kaynak_satiri(sl, c.bolumler.get("kaynak", ""))  # type: ignore[arg-type]
    c.kullanilan.add("kaynak")


def s_uygulanabilirlik_2(sl, c: Icerik, ust: float, kok: Path) -> None:
    y0 = ust + 0.1
    sw = 21.4
    metin(sl, "olcek-baslik", SOL, y0, sw, 0.9, B("Büyüme yolu: arayüz aynı kalır, alt katman değişir"), boyut=16)
    tablo(sl, "olcek", SOL, y0 + 1.0, sw, c.tablo("olcek"), [1.0, 1.25, 1.55], boyut=12.5, satir_h=0.95,
          alt_sinir=y0 + 6.3)
    birim = c.maddeler("birim")
    yb = y0 + 6.6
    kw = (sw - 2 * 0.5) / 3
    for i, madde in enumerate(birim):
        buyuk, aciklama = ayir(madde)
        sayi_karti(sl, f"birim-{i + 1}", SOL + i * (kw + 0.5), yb, kw, 3.2, buyuk, aciklama,
                   renk=LACIVERT, buyuk_boyut=21, boyut=11.5)
    metin(sl, "birim-not", SOL, yb + 3.3, sw, 1.2, c["birim_not"], boyut=11.5, renk=GRI)
    metin(sl, "risk-baslik", SOL, yb + 4.6, sw, 0.9, B("Riskler ve karşı önlemler"), boyut=16)
    tablo(sl, "risk", SOL, yb + 5.5, sw, c.tablo("risk"), [1, 1.75], boyut=12, satir_h=0.9,
          alt_sinir=KAYNAK_UST - 0.2)
    sagx = SOL + sw + 0.9
    saw = SAG - sagx
    kart(sl, "surdurulebilirlik", sagx, y0, saw, KAYNAK_UST - 0.3 - y0, c["surdurulebilirlik"],
         boyut=14.5, sol_serit=YESIL, ic=0.55)
    kaynak_satiri(sl, c.bolumler.get("kaynak", ""))  # type: ignore[arg-type]
    c.kullanilan.add("kaynak")


def s_ozgunluk(sl, c: Icerik, ust: float, kok: Path) -> None:
    metin(sl, "ana", SOL, ust + 0.2, GEN, 1.8, c["ana"], boyut=24, renk=LACIVERT)
    y0 = ust + 2.6
    sw = 15.6
    ozgun = c.maddeler("ozgun")
    hazir = c.maddeler("hazir")
    metin(sl, "katman-baslik", SOL, y0, sw, 0.9, B("Takımın yazdığı katmanlar"), boyut=16)
    y = y0 + 1.1
    oh = 1.55
    for i, madde in enumerate(ozgun):
        kart(sl, f"ozgun-{i + 1}", SOL, y + i * (oh + 0.25), sw, oh, [Blok("paragraf", madde)],
             boyut=13.5, dolgu=LACIVERT, renk=BEYAZ, ic=0.3, dikey=MSO_ANCHOR.MIDDLE)
    y += len(ozgun) * (oh + 0.25) + 0.15
    metin(sl, "hazir-baslik", SOL, y, sw, 0.8, B("Üzerine kurulduğu açık kaynak temel"), boyut=13)
    kart(sl, "hazir", SOL, y + 0.8, sw, KAYNAK_UST - 0.3 - (y + 0.8), [Blok("paragraf", " · ".join(hazir))],
         boyut=13, dolgu=GRI_ACIK, renk=METIN, ic=0.3, hizala=PP_ALIGN.CENTER, dikey=MSO_ANCHOR.MIDDLE)
    sagx = SOL + sw + 0.9
    saw = SAG - sagx
    kw = (saw - 0.6) / 2
    kh = (KAYNAK_UST - 0.3 - y0 - 0.6) / 2
    for i, anahtar in enumerate(("filigran", "sorgu", "literatur", "yerlilik")):
        kart(sl, anahtar, sagx + (i % 2) * (kw + 0.6), y0 + (i // 2) * (kh + 0.6), kw, kh, c[anahtar],
             boyut=15, sol_serit=TURUNCU if anahtar != "yerlilik" else KIRMIZI, ic=0.55)
    kaynak_satiri(sl, c.bolumler.get("kaynak", ""))  # type: ignore[arg-type]
    c.kullanilan.add("kaynak")


def s_hedef(sl, c: Icerik, ust: float, kok: Path) -> None:
    y0 = ust + 0.1
    sw = 26.0
    metin(sl, "aktor-baslik", SOL, y0, sw, 0.9, B("Hedef kullanıcılar ve her birine sağlanan fayda"), boyut=16)
    tablo(sl, "aktorler", SOL, y0 + 1.0, sw, c.tablo("aktorler"), [0.8, 1.35, 1.55], boyut=13.5,
          satir_h=1.9, alt_sinir=y0 + 11.6)
    sagx = SOL + sw + 0.9
    saw = SAG - sagx
    istat = c.maddeler("erisim")
    for i, madde in enumerate(istat):
        buyuk, aciklama = ayir(madde)
        sayi_karti(sl, f"erisim-{i + 1}", sagx, y0 + i * 3.9, saw, 3.5, buyuk, aciklama, renk=MAVI,
                   buyuk_boyut=30, boyut=12)
    y = y0 + len(istat) * 3.9
    kart(sl, "katman", sagx, y, saw, 11.6 - (y - y0), c["katman"], boyut=14, sol_serit=MAVI)
    y2 = y0 + 12.2
    metin(sl, "etki-baslik", SOL, y2, GEN, 0.9, B("Toplumsal ve yöntemsel etki"), boyut=16)
    w = (GEN - 2 * 0.6) / 3
    for i, anahtar in enumerate(("etki_uretici", "etki_remix", "etki_yontem")):
        kart(sl, anahtar, SOL + i * (w + 0.6), y2 + 1.0, w, KAYNAK_UST - 0.3 - (y2 + 1.0), c[anahtar],
             boyut=15, sol_serit=TURUNCU, ic=0.55)
    kaynak_satiri(sl, c.bolumler.get("kaynak", ""))  # type: ignore[arg-type]
    c.kullanilan.add("kaynak")


AYLAR = ["Oca", "Şub", "Mar", "Nis", "May", "Haz", "Tem", "Ağu", "Eyl", "Eki", "Kas", "Ara"]


def tarih(d: date) -> str:
    return f"{d.day} {AYLAR[d.month - 1]} {d.year}"


def s_takvim(sl, c: Icerik, ust: float, kok: Path, paketler=None, bugun: date | None = None) -> None:
    # Sablonun ornek tablosu kaldirilip ayni renklerle 11 paketlik tablo kurulur.
    for sekil in list(sl.shapes):
        if sekil.has_table and not sekil.name.startswith(ONEK):
            sekil._element.getparent().remove(sekil._element)
    faaliyetler = {s[0]: s[1] for s in c.tablo("faaliyetler")[1:]}
    kodlar = [p[0] for p in paketler]
    if list(faaliyetler) != kodlar:
        raise SystemExit(f"14-takvim.md iş paketleri rapor_takvim.PAKETLER ile aynı değil: "
                         f"{list(faaliyetler)} != {kodlar}")
    bugun = bugun or date.today()
    satirlar = [["No", "İş paketi adı", "Alt faaliyetler", "Başlangıç", "Bitiş", "Durum"]]
    for kod, ad, bas, bit, durum in paketler:
        if durum == "bitti":
            etiket = "Tamamlandı"
        elif bas <= bugun:
            etiket = "Sürüyor"
        else:
            etiket = "Planlandı"
        satirlar.append([kod, ad, faaliyetler[kod], tarih(bas), tarih(bit), etiket])
    gf = tablo(sl, "takvim", SOL, ust + 0.1, GEN, satirlar, [0.55, 2.6, 3.3, 1.0, 1.0, 1.0], boyut=12,
               satir_h=(KAYNAK_UST - 0.4 - ust) / len(satirlar), hizalar="llllll", alt_sinir=KAYNAK_UST)
    tbl = gf.table
    for i in range(1, len(satirlar)):
        hucre = tbl.cell(i, 5)
        renk = {"Tamamlandı": YESIL, "Sürüyor": ALTIN, "Planlandı": MAVI}[satirlar[i][5]]
        for run in hucre.text_frame.paragraphs[0].runs:
            run.font.color.rgb = renk
            run.font.bold = True
        for run in tbl.cell(i, 0).text_frame.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = LACIVERT
    kaynak_satiri(sl, c.bolumler.get("kaynak", ""))  # type: ignore[arg-type]
    c.kullanilan.add("kaynak")


YERLESIMLER = {
    "03-ozet.md": s_ozet,
    "04-problem.md": s_problem,
    "05-cozum-1.md": s_cozum_1,
    "06-cozum-2.md": s_cozum_2,
    "07-yontem-1.md": s_yontem_1,
    "08-yontem-2.md": s_yontem_2,
    "09-prototip.md": s_prototip,
    "10-uygulanabilirlik-1.md": s_uygulanabilirlik_1,
    "11-uygulanabilirlik-2.md": s_uygulanabilirlik_2,
    "12-ozgunluk.md": s_ozgunluk,
    "13-hedef-kitle.md": s_hedef,
    "14-takvim.md": s_takvim,
}
